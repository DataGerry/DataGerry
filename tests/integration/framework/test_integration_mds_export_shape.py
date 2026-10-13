# DataGerry - OpenSource Enterprise CMDB
# Copyright (C) 2026 becon GmbH
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.
"""
Integration tests for how each object export format writes a multi-data-section field, against a real MongoDB

``BaseExportWriter.from_database`` reads and renders the seeded object - whose flat field list carries its MDS
fields, as the renderer always gives them - and each format serializes it. Pinned: CSV, XLSX and XML write each MDS
field once, with its row values, and the flat value no row holds appears nowhere; JSON keeps that flat entry next
to the rows, deliberately - it is what the JSON import stores back
"""
import json
import xml.etree.ElementTree as ET
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from cmdb.database import MongoDatabaseManager
from cmdb.framework.exporter.config.exporter_config import ExporterConfig
from cmdb.framework.exporter.format.base_exporter_format import BaseExporterFormat
from cmdb.framework.exporter.format.csv_export_format import CsvExportFormat
from cmdb.framework.exporter.format.json_export_format import JsonExportFormat
from cmdb.framework.exporter.format.xlsx_export_format import XlsxExportFormat
from cmdb.framework.exporter.format.xml_export_format import XmlExportFormat
from cmdb.framework.exporter.writer.base_export_writer import BaseExportWriter
from cmdb.interface.rest_api.responses.response_parameters.collection_parameters import CollectionParameters
from cmdb.models.user_model import CmdbUser
from cmdb.security.acl.permission import AccessControlPermission
from tests.utils import mds_export_seed as seed
# -------------------------------------------------------------------------------------------------------------------- #

FLAT_VALUES: tuple[str, str] = (seed.FLAT_PORT, seed.FLAT_SPEED)


@pytest.fixture(autouse=True)
def _seeded(rest_api, database_manager: MongoDatabaseManager, database_name: str):
    """Seeds the MDS type and object and pushes the app context the renderer resolves its managers through"""
    seed.seed(database_manager, database_name)

    with rest_api.application.test_request_context():
        yield

    seed.purge(database_manager, database_name)


def _export(export_format: BaseExporterFormat, dbm: MongoDatabaseManager, user: CmdbUser) -> bytes:
    """The real export of the seeded object"""
    parameters = CollectionParameters(filter={'type_id': seed.TYPE_ID})
    writer = BaseExportWriter(export_format, ExporterConfig(parameters=parameters, options={}))
    writer.from_database(dbm, user, AccessControlPermission.READ)

    return writer.export().get_data()


def _row_values() -> list[str]:
    """Every value the rows hold"""
    return [value for row in seed.stored_rows() for _, value in row]


class TestTabularFormats:
    """CSV and XLSX: one column per MDS field, one line per row."""

    def test_csv(self, database_manager, full_access_user: CmdbUser) -> None:
        """The header names each MDS field once; the rows carry their values; no flat value"""
        lines = _export(CsvExportFormat(), database_manager, full_access_user).decode('utf-8').splitlines()

        header = lines[0].split(',')
        assert [header.count(name) for name in seed.MDS_FIELDS] == [1, 1]
        assert [line.split(',')[-2:] for line in lines[1:]] == [list(row) for row in seed.ROWS]
        assert not any(flat in '\n'.join(lines) for flat in FLAT_VALUES)

    def test_xlsx(self, database_manager, full_access_user: CmdbUser) -> None:
        """The same on the type's sheet"""
        sheet = load_workbook(BytesIO(_export(XlsxExportFormat(), database_manager, full_access_user))).active
        rows: list[tuple[Any, ...]] = list(sheet.iter_rows(values_only=True))

        assert [rows[0].count(name) for name in seed.MDS_FIELDS] == [1, 1]
        assert [list(row[-2:]) for row in rows[1:]] == [list(row) for row in seed.ROWS]
        assert not any(flat in cell for row in rows for cell in row if isinstance(cell, str) for flat in FLAT_VALUES)


class TestXml:
    """XML: an MDS field in its rows only."""

    def test_fields_hold_the_regular_fields_and_the_rows_the_mds_ones(
            self, database_manager, full_access_user: CmdbUser) -> None:
        """<fields> without the MDS names; each row with its own values; the flat value nowhere"""
        document: bytes = _export(XmlExportFormat(), database_manager, full_access_user)
        obj = ET.fromstring(document).find('object')

        assert [field.attrib['name'] for field in obj.findall('fields/field')] == [seed.NAME_FIELD]
        rows = obj.findall('multi_data_sections/section/row')
        assert [[(f.attrib['name'], f.attrib['value']) for f in row.findall('field')] for row in rows] == \
            seed.stored_rows()
        assert not any(flat.encode() in document for flat in FLAT_VALUES)


class TestJson:
    """JSON: the flat entry kept next to the rows - the import stores both back."""

    def test_the_flat_entries_and_the_rows(self, database_manager, full_access_user: CmdbUser) -> None:
        """`fields` is the stored flat list, MDS entries included; the rows are the stored rows"""
        exported = json.loads(_export(JsonExportFormat(), database_manager, full_access_user))[0]

        assert exported['fields'] == seed.stored_flat_fields()
        rows = exported['multi_data_sections'][0]['values']
        assert [[(d['name'], d['value']) for d in row['data']] for row in rows] == seed.stored_rows()
        assert all(value not in [f['value'] for f in exported['fields']] for value in _row_values())
