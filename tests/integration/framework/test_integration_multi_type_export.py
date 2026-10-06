# DataGerry - OpenSource Enterprise CMDB
# Copyright (C) 2026 becon GmbH
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.
"""
Integration tests for an object export spanning two types, against a real MongoDB and the real renderer

``BaseExportWriter.from_database`` reads and renders the selection, the format serializes it. Pinned:

  - the rendered objects of two types carry each type's own fields, which is what the XML and the XLSX
    export now lay out per object / per sheet - nothing of the other type leaks in
  - a ZIP packs one file per type, each holding that type's fields only and each the direct export of that
    type with the same options (here a human-readable CSV, run through the writer's location resolution)
  - the same export twice gives the same bytes (re-run safe)
"""
import xml.etree.ElementTree as ET
import zipfile
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from cmdb.database import MongoDatabaseManager
from cmdb.framework.exporter.config.exporter_config import ExporterConfig
from cmdb.framework.exporter.format.base_exporter_format import BaseExporterFormat
from cmdb.framework.exporter.format.csv_export_format import CsvExportFormat
from cmdb.framework.exporter.format.xlsx_export_format import XlsxExportFormat
from cmdb.framework.exporter.format.xml_export_format import XmlExportFormat
from cmdb.framework.exporter.format.zip_export_format import ZipExportFormat
from cmdb.framework.exporter.writer.base_export_writer import BaseExportWriter
from cmdb.interface.rest_api.responses.response_parameters.collection_parameters import CollectionParameters
from cmdb.models.user_model import CmdbUser
from cmdb.security.acl.permission import AccessControlPermission

from tests.utils import multi_type_export_seed as seed
from tests.utils.multi_type_export_seed import (
    ROUTER_FIELDS,
    ROUTER_LABEL,
    ROUTER_OBJECT_ID,
    ROUTER_VALUES,
    SEEDED_TYPE_IDS,
    SERVER_FIELDS,
    SERVER_LABEL,
    SERVER_OBJECT_ID,
    SERVER_VALUES,
)
# -------------------------------------------------------------------------------------------------------------------- #

XML_FORMAT: str = 'XmlExportFormat'
CSV_FORMAT: str = 'CsvExportFormat'

# What a ZIP of the two-type selection holds: one file per type, named by the type
ZIP_MEMBER_COUNT: int = 2


@pytest.fixture(autouse=True)
def _seeded(rest_api, database_manager: MongoDatabaseManager, database_name: str):
    """Seeds the two types and pushes the app context the renderer resolves its managers through"""
    seed.purge(database_manager, database_name)
    seed.seed(database_manager, database_name)

    with rest_api.application.test_request_context():
        yield

    seed.purge(database_manager, database_name)


def _writer(export_format: BaseExporterFormat, dbm: MongoDatabaseManager, user: CmdbUser,
            options: dict[str, Any] | None = None, type_ids: list[int] | None = None) -> BaseExportWriter:
    """A writer holding the rendered two-type selection, sorted by public_id"""
    parameters = CollectionParameters(filter={'type_id': {'$in': type_ids or SEEDED_TYPE_IDS}}, sort='public_id',
                                      order=1)
    writer = BaseExportWriter(export_format, ExporterConfig(parameters=parameters, options=options or {}))
    writer.from_database(dbm, user, AccessControlPermission.READ)

    return writer


def _xml_fields(document: str | bytes) -> dict[str, list[tuple[str, str]]]:
    """The (name, value) field pairs of each exported <object>, keyed by its public_id"""
    return {
        obj.find('meta/public_id').text: [(field.attrib['name'], field.attrib['value']) for field in obj.find('fields')]
        for obj in ET.fromstring(document).findall('object')
    }


class TestRenderedSelection:
    """What the formats are handed: rendered objects, each with its own type's fields"""

    def test_each_rendered_object_carries_its_own_types_fields(self, database_manager, full_access_user) -> None:
        """The renderer builds the field list from the object's type - the formats rely on it"""
        writer = _writer(XmlExportFormat(), database_manager, full_access_user)

        assert [[field['name'] for field in obj.fields] for obj in writer.data] == [SERVER_FIELDS, ROUTER_FIELDS]


class TestXmlExport:
    """The real XML export of the two-type selection"""

    def test_each_object_carries_its_own_types_fields(self, database_manager, full_access_user) -> None:
        """No field of the other type appears, each type keeps its own field order"""
        response = _writer(XmlExportFormat(), database_manager, full_access_user).export()

        assert _xml_fields(response.get_data()) == {
            str(SERVER_OBJECT_ID): list(SERVER_VALUES.items()),
            str(ROUTER_OBJECT_ID): list(ROUTER_VALUES.items()),
        }

    def test_the_same_export_twice_gives_the_same_bytes(self, database_manager, full_access_user) -> None:
        """Re-run safe: nothing in the document depends on the run"""
        first = _writer(XmlExportFormat(), database_manager, full_access_user).export().get_data()
        second = _writer(XmlExportFormat(), database_manager, full_access_user).export().get_data()

        assert first == second


class TestXlsxExport:
    """The real XLSX export of the two-type selection"""

    def test_one_sheet_per_type_with_its_own_columns_and_values(self, database_manager, full_access_user) -> None:
        """Each sheet: the type's own header and its one object's row"""
        response = _writer(XlsxExportFormat(), database_manager, full_access_user).export()
        workbook = load_workbook(BytesIO(response.get_data()))

        assert workbook.sheetnames == [SERVER_LABEL, ROUTER_LABEL]

        expected_sheets = (
            (SERVER_LABEL, SERVER_FIELDS, SERVER_OBJECT_ID, SERVER_VALUES),
            (ROUTER_LABEL, ROUTER_FIELDS, ROUTER_OBJECT_ID, ROUTER_VALUES),
        )

        for label, field_names, object_id, values in expected_sheets:
            sheet = workbook[label]
            assert [cell.value for cell in sheet[1]] == ['public_id', 'active', *field_names]
            assert [cell.value for cell in sheet[2]] == [str(object_id), 'True', *values.values()]


class TestZipExport:
    """The ZIP wrapper packs one file per type"""

    def test_each_packed_file_holds_its_types_fields_only(self, database_manager, full_access_user) -> None:
        """Two XML files, each with one object and that object's own fields"""
        options = {'classname': XML_FORMAT}
        response = _writer(ZipExportFormat(), database_manager, full_access_user, options).export()

        with zipfile.ZipFile(BytesIO(response.get_data())) as archive:
            documents = [archive.read(name) for name in archive.namelist()]

        assert len(documents) == ZIP_MEMBER_COUNT
        packed: dict[str, list[tuple[str, str]]] = {}
        for document in documents:
            packed.update(_xml_fields(document))

        assert packed == {
            str(SERVER_OBJECT_ID): list(SERVER_VALUES.items()),
            str(ROUTER_OBJECT_ID): list(ROUTER_VALUES.items()),
        }

    def test_a_human_readable_zip_packs_the_direct_human_readable_exports(
        self, database_manager, full_access_user,
    ) -> None:
        """The flag reaches the packed CSVs, and the archive's name says readable"""
        options = {'classname': CSV_FORMAT, 'human_readable': 'true'}
        response = _writer(ZipExportFormat(), database_manager, full_access_user, options).export()

        with zipfile.ZipFile(BytesIO(response.get_data())) as archive:
            packed = sorted(archive.read(name) for name in archive.namelist())

        direct = sorted(
            _writer(CsvExportFormat(), database_manager, full_access_user, options, [type_id]).export().get_data()
            for type_id in SEEDED_TYPE_IDS
        )

        assert packed == direct
        assert '_readable.zip' in response.headers['Content-Disposition']
