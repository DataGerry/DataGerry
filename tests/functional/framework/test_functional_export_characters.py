# DataGerry - OpenSource Enterprise CMDB
# Copyright (C) 2026 becon GmbH
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.
"""
What the object export route answers for characters an export format cannot write as they are - end to end

* a line feed, carriage return or tab in a field value: every format writes it so that a reader gets it back - the
  XML as a character reference, since XML would read a raw one in an attribute as a space
* a control character XML 1.0 has no spelling for: XML and XLSX (XML inside) refuse the export with a 400 that
  names the object and the field - XLSX used to answer it with a 500; CSV and JSON write it
"""
import json
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from http import HTTPStatus
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from cmdb.database import MongoDatabaseManager
from cmdb.models.object_model import CmdbObject
from cmdb.models.type_model import CmdbType
from tests.utils.ipam_doc_builders import make_type_doc
# -------------------------------------------------------------------------------------------------------------------- #

EXPORT_URL: str = '/exporter/'
TYPE_ID: int = 98301
OBJECT_ID: int = 98311
NOTE_FIELD: str = 'note'
MULTI_LINE: str = 'line1\nline2\r\nthird\ttab'
BELL: str = 'bell\x07'
REFUSAL: str = f"Object {OBJECT_ID}, field '{NOTE_FIELD}'"


@pytest.fixture(name='seed_note')
def fixture_seed_note(database_manager: MongoDatabaseManager, database_name: str):
    """Seeds the type; answers a function writing its one object with a given note"""
    types = database_manager.get_collection(CmdbType.COLLECTION, database_name)
    objects = database_manager.get_collection(CmdbObject.COLLECTION, database_name)
    types.delete_many({'public_id': TYPE_ID})
    types.insert_one(make_type_doc(TYPE_ID, 'export-characters',
                                   fields=[{'type': 'text', 'name': NOTE_FIELD, 'label': 'Note'}]))

    def _seed(note: str) -> None:
        objects.delete_many({'type_id': TYPE_ID})
        objects.insert_one({'public_id': OBJECT_ID, 'type_id': TYPE_ID, 'active': True, 'author_id': 1,
                            'version': '1.0.0', 'creation_time': datetime.now(timezone.utc),
                            'fields': [{'name': NOTE_FIELD, 'value': note}]})

    yield _seed

    types.delete_many({'public_id': TYPE_ID})
    objects.delete_many({'type_id': TYPE_ID})


def _export(rest_api, classname: str, **params: Any):
    """The export route for the seeded type"""
    return rest_api.get(EXPORT_URL, query_string={'filter': json.dumps({'type_id': TYPE_ID}), 'classname': classname,
                                                  **params})


class TestWhitespaceReadsBack:
    """A multi-line value is what a reader of the file gets."""

    def test_xml(self, rest_api, seed_note) -> None:
        """As character references - parsed back unchanged"""
        seed_note(MULTI_LINE)

        response = _export(rest_api, 'XmlExportFormat')

        assert response.status_code == HTTPStatus.OK
        assert ET.fromstring(response.data).find('object/fields/field').attrib['value'] == MULTI_LINE

    def test_xlsx(self, rest_api, seed_note) -> None:
        """The cell holds it as it is"""
        seed_note(MULTI_LINE)

        sheet = load_workbook(BytesIO(_export(rest_api, 'XlsxExportFormat').data)).active

        assert sheet.cell(row=2, column=sheet.max_column).value == MULTI_LINE


class TestAnUnwritableCharacter:
    """A control character XML 1.0 cannot carry."""

    @pytest.mark.parametrize('classname', ['XmlExportFormat', 'XlsxExportFormat'])
    def test_xml_and_xlsx_answer_400_naming_object_and_field(self, rest_api, seed_note, classname: str) -> None:
        """Not a 500 - and the reader can find the value"""
        seed_note(BELL)

        response = _export(rest_api, classname)

        assert response.status_code == HTTPStatus.BAD_REQUEST
        assert REFUSAL in response.get_json()['message']
        assert 'U+0007' in response.get_json()['message']

    def test_a_zip_of_xlsx_answers_the_same_400(self, rest_api, seed_note) -> None:
        """The packed format refuses inside the ZIP the same way"""
        seed_note(BELL)

        response = _export(rest_api, 'XlsxExportFormat', zip='true')

        assert response.status_code == HTTPStatus.BAD_REQUEST
        assert REFUSAL in response.get_json()['message']

    @pytest.mark.parametrize('classname', ['CsvExportFormat', 'JsonExportFormat'])
    def test_csv_and_json_write_it(self, rest_api, seed_note, classname: str) -> None:
        """They can carry it - the formats a refusal points to"""
        seed_note(BELL)

        assert _export(rest_api, classname).status_code == HTTPStatus.OK
