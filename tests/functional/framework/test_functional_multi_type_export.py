# DataGerry - OpenSource Enterprise CMDB
# Copyright (C) 2026 becon GmbH
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.
"""
Functional coverage for an object export whose selection spans several types, through GET /exporter/

Two types sharing one field name (in a different position) are seeded, one object each. Pinned:

  - XML: one document, each <object> carrying its own type's fields only, in its type's order
  - XLSX: one sheet per type with that type's columns; two types sharing a label get two sheets
  - a render-view column selection is narrowed to each type's fields, a missing `columns` key is the default
  - a header entry that is no XML element name, and a value XML cannot carry, are a 400 - not a 500
  - ZIP: each packed file is the direct export of its type with the same options (view, selection, flag)
  - the `readable` filename marker is set only when the format honours the flag
"""
import json
import xml.etree.ElementTree as ET
import zipfile
from http import HTTPStatus
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from cmdb.database import MongoDatabaseManager
from cmdb.models.type_model import CmdbType
from cmdb.models.object_model import CmdbObject

from tests.utils import multi_type_export_seed as seed
from tests.utils.multi_type_export_seed import (
    ROUTER_FIELD,
    ROUTER_FIELDS,
    ROUTER_LABEL,
    ROUTER_OBJECT_ID,
    ROUTER_TYPE_ID,
    ROUTER_VALUES,
    SEEDED_TYPE_IDS,
    SERVER_FIELD,
    SERVER_FIELDS,
    SERVER_LABEL,
    SERVER_OBJECT_ID,
    SERVER_TYPE_ID,
    SERVER_VALUES,
    SHARED_FIELD,
)
# -------------------------------------------------------------------------------------------------------------------- #

EXPORT_URL: str = '/exporter/'
XML_FORMAT: str = 'XmlExportFormat'
XLSX_FORMAT: str = 'XlsxExportFormat'
CSV_FORMAT: str = 'CsvExportFormat'
JSON_FORMAT: str = 'JsonExportFormat'

# The `readable` marker closing the stem of a presentation export's filename
READABLE_FILENAME_SUFFIX: str = '_readable.'

# An object value holding a character XML 1.0 cannot carry (a vertical tab)
UNCARRIABLE_VALUE: str = 'bad\x0bvalue'

SELECTION_FILTER: str = json.dumps({'type_id': {'$in': SEEDED_TYPE_IDS}})


@pytest.fixture(name='collections')
def fixture_collections(database_manager: MongoDatabaseManager, database_name: str):
    """The type and object collections, purged of the seeded documents before and after"""
    seed.purge(database_manager, database_name)
    yield (database_manager.get_collection(CmdbType.COLLECTION, database_name),
           database_manager.get_collection(CmdbObject.COLLECTION, database_name))
    seed.purge(database_manager, database_name)


@pytest.fixture(name='seeded')
def fixture_seeded(collections, database_manager: MongoDatabaseManager, database_name: str):
    """A Server type and a Router type, one object each"""
    seed.seed(database_manager, database_name)

    return collections


def _export(rest_api, export_format: str, **params: str):
    """Exports the seeded selection in the given format"""
    query = {'filter': SELECTION_FILTER, 'classname': export_format, 'sort': 'public_id', 'order': '1', **params}

    return rest_api.get(EXPORT_URL, query_string=query)


def _render_params(columns: list[str] | None, header: list[str] | None = None) -> dict[str, str]:
    """The query parameters of a render-view export selecting the given header and columns"""
    metadata: dict[str, Any] = {'header': header or ['public_id']}

    if columns is not None:
        metadata['columns'] = columns

    return {'view': 'render', 'metadata': json.dumps(metadata)}


def _xml_fields(response) -> dict[str, list[tuple[str, str]]]:
    """The (name, value) field pairs of each exported <object>, keyed by its public_id"""
    root = ET.fromstring(response.get_data(as_text=True))

    return {
        obj.find('meta/public_id').text: [(field.attrib['name'], field.attrib['value']) for field in obj.find('fields')]
        for obj in root.findall('object')
    }


def _header_row(sheet) -> list[Any]:
    """The non-empty cells of a worksheet's header row"""
    return [cell.value for cell in sheet[1] if cell.value is not None]


class TestMultiTypeXml:
    """One XML document for the whole selection, each object with its own type's fields"""

    @pytest.mark.usefixtures('seeded')
    def test_each_object_carries_its_own_types_fields(self, rest_api) -> None:
        """No field of the other type appears, and each type keeps its own field order"""
        response = _export(rest_api, XML_FORMAT)

        assert response.status_code == HTTPStatus.OK
        assert _xml_fields(response) == {
            str(SERVER_OBJECT_ID): list(SERVER_VALUES.items()),
            str(ROUTER_OBJECT_ID): list(ROUTER_VALUES.items()),
        }

    @pytest.mark.usefixtures('seeded')
    def test_render_selection_is_narrowed_to_each_type(self, rest_api) -> None:
        """Each object gets the selected fields its type owns, in the order of the selection"""
        response = _export(rest_api, XML_FORMAT, **_render_params([ROUTER_FIELD, SERVER_FIELD, SHARED_FIELD]))

        assert response.status_code == HTTPStatus.OK
        assert _xml_fields(response) == {
            str(SERVER_OBJECT_ID): [(SERVER_FIELD, SERVER_VALUES[SERVER_FIELD]),
                                    (SHARED_FIELD, SERVER_VALUES[SHARED_FIELD])],
            str(ROUTER_OBJECT_ID): [(ROUTER_FIELD, ROUTER_VALUES[ROUTER_FIELD]),
                                    (SHARED_FIELD, ROUTER_VALUES[SHARED_FIELD])],
        }

    @pytest.mark.usefixtures('seeded')
    def test_metadata_without_columns_keeps_every_own_field(self, rest_api) -> None:
        """A missing `columns` key is the default, not an empty selection"""
        response = _export(rest_api, XML_FORMAT, **_render_params(None))

        assert response.status_code == HTTPStatus.OK
        assert _xml_fields(response)[str(SERVER_OBJECT_ID)] == list(SERVER_VALUES.items())

    @pytest.mark.usefixtures('seeded')
    @pytest.mark.parametrize('head', ['bad name', '1st', 'a:b'])
    def test_a_header_that_is_no_element_name_is_a_400(self, rest_api, head: str) -> None:
        """It answered 500: the entry became a tag the document could not be read back with"""
        response = _export(rest_api, XML_FORMAT, **_render_params([], header=['public_id', head]))

        assert response.status_code == HTTPStatus.BAD_REQUEST
        assert 'XML element name' in response.get_json()['message']

    def test_a_value_xml_cannot_carry_is_a_400(self, rest_api, collections) -> None:
        """It answered 500 - now the refusal names the character"""
        types, objects = collections
        types.insert_one(seed.type_doc(SERVER_TYPE_ID, SERVER_LABEL, SERVER_FIELDS))
        objects.insert_one(seed.object_doc(SERVER_OBJECT_ID, SERVER_TYPE_ID, {SHARED_FIELD: UNCARRIABLE_VALUE}))

        response = _export(rest_api, XML_FORMAT)

        assert response.status_code == HTTPStatus.BAD_REQUEST
        assert 'U+000B' in response.get_json()['message']


class TestMultiTypeXlsx:
    """One worksheet per type, each with its own type's columns"""

    @pytest.mark.usefixtures('seeded')
    def test_each_sheet_carries_its_types_columns(self, rest_api) -> None:
        """Sheets in type-id order, each header the identity columns plus that type's fields"""
        response = _export(rest_api, XLSX_FORMAT)
        workbook = load_workbook(BytesIO(response.data))

        assert response.status_code == HTTPStatus.OK
        assert workbook.sheetnames == [SERVER_LABEL, ROUTER_LABEL]
        assert _header_row(workbook[SERVER_LABEL]) == ['public_id', 'active', *SERVER_FIELDS]
        assert _header_row(workbook[ROUTER_LABEL]) == ['public_id', 'active', *ROUTER_FIELDS]

    @pytest.mark.usefixtures('seeded')
    def test_render_selection_is_narrowed_to_each_sheet(self, rest_api) -> None:
        """Each sheet carries the selected fields its type owns, in the order of the selection"""
        response = _export(rest_api, XLSX_FORMAT, **_render_params([ROUTER_FIELD, SERVER_FIELD, SHARED_FIELD]))
        workbook = load_workbook(BytesIO(response.data))

        assert _header_row(workbook[SERVER_LABEL]) == ['public_id', SERVER_FIELD, SHARED_FIELD]
        assert _header_row(workbook[ROUTER_LABEL]) == ['public_id', ROUTER_FIELD, SHARED_FIELD]

    @pytest.mark.usefixtures('collections')
    def test_two_types_sharing_a_label_get_two_sheets(self, rest_api, database_manager, database_name) -> None:
        """The second sheet is numbered instead of renamed by the spreadsheet library"""
        seed.seed(database_manager, database_name, router_label=SERVER_LABEL)

        workbook = load_workbook(BytesIO(_export(rest_api, XLSX_FORMAT).data))

        assert workbook.sheetnames == [SERVER_LABEL, f'{SERVER_LABEL}~2']
        assert _header_row(workbook[f'{SERVER_LABEL}~2']) == ['public_id', 'active', *ROUTER_FIELDS]


def _zip_members(response) -> dict[str, bytes]:
    """The packed files of a ZIP export, by entry name"""
    with zipfile.ZipFile(BytesIO(response.data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _direct_export(rest_api, type_id: int, export_format: str, **params: str) -> bytes:
    """The direct export of one seeded type, with the same options"""
    query = {'filter': json.dumps({'type_id': type_id}), 'classname': export_format, 'sort': 'public_id',
             'order': '1', **params}

    return rest_api.get(EXPORT_URL, query_string=query).data


def _entry_name(type_id: int, extension: str) -> str:
    """The ZIP entry of one seeded type"""
    return f'multi-export-{type_id}_ID_{type_id}.{extension}'


@pytest.mark.usefixtures('seeded')
class TestZipExport:
    """A ZIP packs one file per type, each the direct export of that type with the request's options"""

    @pytest.mark.parametrize('params', [
        {},
        _render_params([ROUTER_FIELD, SERVER_FIELD, SHARED_FIELD]),
        {'human_readable': 'true'},
    ])
    def test_each_packed_csv_is_the_direct_export_of_its_type(self, rest_api, params: dict[str, str]) -> None:
        """Native, a render selection, the human-readable flag: all reach the packed files"""
        members = _zip_members(_export(rest_api, CSV_FORMAT, zip='true', **params))

        assert members == {
            _entry_name(type_id, 'csv'): _direct_export(rest_api, type_id, CSV_FORMAT, **params)
            for type_id in (SERVER_TYPE_ID, ROUTER_TYPE_ID)
        }

    def test_the_frontends_call_shape_stays_a_native_export(self, rest_api) -> None:
        """The reference table zips with no view and the metadata text 'undefined'"""
        members = _zip_members(_export(rest_api, XML_FORMAT, zip='true', metadata='undefined'))

        assert members[_entry_name(SERVER_TYPE_ID, 'xml')] == _direct_export(rest_api, SERVER_TYPE_ID, XML_FORMAT)

    def test_a_render_selection_is_narrowed_inside_the_archive(self, rest_api) -> None:
        """The packed CSV header is the identity column plus the selected fields its type owns"""
        members = _zip_members(_export(rest_api, CSV_FORMAT, zip='true', **_render_params([SERVER_FIELD])))

        assert members[_entry_name(SERVER_TYPE_ID, 'csv')].decode().splitlines()[0] == f'public_id,{SERVER_FIELD}'
        assert members[_entry_name(ROUTER_TYPE_ID, 'csv')].decode().splitlines()[0] == 'public_id'


@pytest.mark.usefixtures('seeded')
class TestReadableMarker:
    """The filename says `readable` only when the file really is a presentation export"""

    @pytest.mark.parametrize(('export_format', 'zipped', 'marked'), [
        (CSV_FORMAT, False, True),
        (XLSX_FORMAT, False, True),
        (JSON_FORMAT, False, False),
        (XML_FORMAT, False, False),
        (CSV_FORMAT, True, True),
        (JSON_FORMAT, True, False),
    ])
    def test_the_marker_follows_the_format(self, rest_api, export_format: str, zipped: bool, marked: bool) -> None:
        """JSON and XML write the same file with the flag - their name used to say readable anyway"""
        params = {'human_readable': 'true', **({'zip': 'true'} if zipped else {})}
        # One type: a CSV refuses the mixed selection
        response = rest_api.get(EXPORT_URL, query_string={
            'filter': json.dumps({'type_id': SERVER_TYPE_ID}), 'classname': export_format, **params,
        })

        assert response.status_code == HTTPStatus.OK
        assert (READABLE_FILENAME_SUFFIX in response.headers['Content-Disposition']) is marked
